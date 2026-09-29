"""Create the project pitch deck and the detailed programming guide.

Run with: .\\.venv\\Scripts\\python.exe scripts/generate_project_documents.py
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables"

NAVY = "102A43"
BLUE = "176B87"
TEAL = "0B8F8C"
ORANGE = "F4A261"
INK = "18324A"
MUTED = "5E7184"
LIGHT = "F5F8FA"
WHITE = "FFFFFF"
LINE = "D8E2EA"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def ppt_rgb(value: str) -> PptRGBColor:
    return PptRGBColor.from_string(value)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Pi-OT Security Probe | Technical Programming Guide | ")
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, color in (("Title", 30, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 12, BLUE)):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = rgb(NAVY)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.space_after = Pt(5)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        add_page_number(footer)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = rgb(WHITE)
            run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def add_code(document: Document, code: str) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.add_run(code)


def new_page(document: Document, number: int, title: str, kicker: str) -> None:
    if number > 1:
        document.add_page_break()
    document.add_paragraph(kicker.upper(), style="Subtitle")
    document.paragraphs[-1].runs[0].font.color.rgb = rgb(TEAL)
    document.paragraphs[-1].runs[0].font.size = Pt(9)
    document.add_heading(title, level=1)


PAGES = [
    {
        "title": "Pi-OT Security Probe",
        "kicker": "Technical programming guide | page 1 of 10",
        "paragraphs": [
            "Pi-OT Security Probe is a local, offline-first engineering console for collecting and reviewing bounded network evidence. The working software runs on Windows, Linux, or Kali without Raspberry Pi hardware. It is designed for an authorized operator who needs a readable inventory of devices, selected service observations, web metadata, Wi-Fi access-point information, and changes over time.",
            "The program intentionally reports observations rather than inferred vulnerabilities, risk scores, or guessed device types. This keeps the output traceable: every displayed fact has a collector, a scan record, a timestamp, and stored evidence behind it.",
        ],
        "bullets": [
            "Primary interface: a FastAPI website bound to 127.0.0.1:8088 by the root run.ps1 or run.sh launcher.",
            "Main live features: attached-network discovery, confirmed single-host TCP checks, one-request HTTP/HTTPS inspection, Wi-Fi inventory, baselines, change review, and JSON/CSV exports.",
            "Development goal: code is organised by responsibility under src/pi_ot_probe so an edit has one obvious home.",
            "Non-goal: exploit testing, password testing, broad port scanning, operating-system detection, NSE scripts, packet capture, and production OT protocol interaction.",
        ],
        "table": (
            ["Audience", "What they use it for"],
            [
                ["Local operator", "Discover and review real evidence on an authorized private network."],
                ["Developer", "Change one layer at a time and validate it with fixtures plus local live checks."],
                ["Reviewer", "Trace a UI observation back to stored scan evidence and a collector."],
            ],
        ),
    },
    {
        "title": "Runtime and source layout",
        "kicker": "Page 2 of 10 | where code lives",
        "paragraphs": [
            "The project uses a Python src layout. The package is installed in editable mode during development, so source changes are used immediately after the server restarts. The root launchers create a virtual environment, install the project, and start uvicorn with pi_ot_probe.web.app:app.",
            "The earlier kali_demo folder is retained only for compatibility launchers and existing local SQLite data. New work belongs in src/pi_ot_probe. This separation removes the ambiguity between the runnable web application, hardware ideas, simulation support, and unfinished experiments.",
        ],
        "table": (
            ["Folder", "Responsibility"],
            [
                ["web/", "Application assembly, HTTP routes, request schemas, security headers, HTML, CSS, JavaScript views."],
                ["application/", "Use cases: scan coordination, baselines, dashboard state, validation-lab descriptions."],
                ["collectors/", "Adapters for interfaces, target policy, subprocess execution, Nmap, Wi-Fi, and HTTP/TLS."],
                ["domain/", "Shared typed records: Asset, Service, Scan, protocol and Wi-Fi observations."],
                ["database/", "SQLite schema, repository methods, and current-asset reconstruction."],
                ["scanners/", "Scanner contract, cancellation token, lifecycle orchestration, progress, and audit logging."],
                ["reports/ and diagnostics/", "Raw JSON/CSV exports and a repeatable local-machine verification command."],
                ["experimental/ and hardware/", "Deferred work; isolated so it cannot blur live website behavior."],
            ],
        ),
        "bullets": [
            "docs/development.md is the learning map: it names the file to change for a UI label, route, collector, baseline rule, or report.",
            "tests/web contains feature tests for the website, collectors, baselines, reports, and live loopback HTTP assembly.",
        ],
    },
    {
        "title": "How a browser action becomes stored evidence",
        "kicker": "Page 3 of 10 | request flow",
        "paragraphs": [
            "The frontend deliberately separates control, API calls, DOM helpers, and visual views. web/static/js/app.js owns button events and refresh behavior. api.js sends requests and turns failed responses into useful errors. The view modules render data with textContent, so network-provided strings are never treated as HTML.",
            "FastAPI receives the request in a small route group. A Pydantic request schema validates the body. The route calls ProbeService rather than performing a scan directly. This keeps HTTP status handling separate from operational logic and allows the service to be tested without a browser.",
        ],
        "table": (
            ["Step", "Code path", "Purpose"],
            [
                ["1", "web/static/js/app.js → api.js", "Read UI input and send a JSON request."],
                ["2", "web/routes/scans.py + web/schemas.py", "Validate input and translate exceptions to HTTP responses."],
                ["3", "application/service.py", "Confirm authorization, serialize operations, choose a collector, track progress."],
                ["4", "collectors/*.py", "Read local state or run a bounded external tool."],
                ["5", "scanners/orchestrator.py", "Persist scan state, snapshots, audit events, cancellation, and errors."],
                ["6", "application/dashboard.py", "Build score-free current state from immutable observations."],
                ["7", "web/static/js/views/*.js", "Render inventory, changes, history, Wi-Fi, and validation information."],
            ],
        ),
        "code": "POST /api/real/run  →  DiscoveryRequest  →  ProbeService.run_real()\n→ validate_target()  →  NmapDiscoveryScanner  →  ScanOrchestrator.run()\n→ SQLite scan_assets snapshot  →  GET /api/state  →  renderAssets()",
    },
    {
        "title": "How local networks are discovered and approved",
        "kicker": "Page 4 of 10 | discovery policy",
        "paragraphs": [
            "Network discovery starts by reading the local computer, not by probing the network. collectors/interfaces.py prefers psutil to enumerate active IPv4 interfaces. On Windows it can fall back to a PowerShell adapter; on Linux it can fall back to iproute2 JSON. Loopback and non-private addresses are excluded from discoverable targets.",
            "Each active RFC1918 interface becomes a LocalNetwork record containing the interface name, local address, full connected network, and a safe target. A broad private interface such as 10.0.0.0/8 is never exposed as an unrestricted scan. The UI offers a target limited to a maximum of 256 addresses, normally a /24.",
        ],
        "table": (
            ["Validation rule", "Why it exists"],
            [
                ["IPv4 only", "Keeps the first collector and target model explicit and bounded."],
                ["RFC1918 only", "Rejects public and special-use ranges."],
                ["Canonical network required", "Rejects ambiguous forms such as a host address with a subnet suffix."],
                ["Maximum 256 addresses", "Prevents accidental broad discovery; /24 or smaller only."],
                ["Must be a subnet of an active interface", "Prevents scanning an unrelated private network."],
                ["Explicit authorization flag", "The server rejects discovery before any Nmap process is created."],
            ],
        ),
        "code": "validate_target('192.168.100.0/24', local_networks)\n# Accepts only an attached RFC1918 IPv4 network with <= 256 addresses.",
        "bullets": [
            "The website asks the operator to select a local target; it does not accept arbitrary Internet targets.",
            "The target validation is on the server, so modifying browser JavaScript cannot bypass it.",
        ],
    },
    {
        "title": "Bounded collectors: Nmap, web, and Wi-Fi",
        "kicker": "Page 5 of 10 | what real collectors do",
        "paragraphs": [
            "The discovery collector runs Nmap host discovery only: -sn, a limited maximum rate, one retry, a host timeout, XML output, and a three-minute outer timeout. XML is parsed into Asset records only when Nmap says the host is up. The parser retains IP, MAC, vendor, hostnames, discovery reason, source, and timestamp when observed.",
            "Service verification is a separate, explicitly confirmed operation for one asset from the current discovery. It runs a TCP connect scan against Nmap's 20 common ports, with a maximum rate of five connections per second, one retry, a per-host timeout, and an outer one-minute timeout. It records only open TCP ports and Nmap's structured service details.",
        ],
        "table": (
            ["Collector", "Network action", "Evidence retained"],
            [
                ["Host discovery", "Nmap -sn on an approved /24-or-smaller subnet.", "Up hosts, IP/MAC/vendor/hostnames, Nmap reason."],
                ["Service verification", "Nmap TCP connect scan of 20 ports on one selected host.", "Open ports, service name, product/version fields when Nmap returns them."],
                ["HTTP/HTTPS inspection", "One HEAD / request to an already observed web port; no redirect following.", "Status, headers, HTTP version, TLS details, certificate fingerprint; Set-Cookie values are redacted."],
                ["Wi-Fi inventory", "Normal OS-managed scan through nmcli or netsh after confirmation.", "SSID/BSSID, signal, channel, encryption and other OS-provided radio fields."],
            ],
        ),
        "bullets": [
            "collectors/process.py uses argument arrays, captures output, and kills a timed-out subprocess before raising an error.",
            "No collector runs OS detection, version scans as a purpose, NSE scripts, credential tests, exploit code, deauthentication, monitor mode, or packet capture.",
        ],
    },
    {
        "title": "Scan lifecycle, cancellation, and evidence integrity",
        "kicker": "Page 6 of 10 | scanner orchestration",
        "paragraphs": [
            "Every collector implements the Scanner contract and receives a ScanContext plus a shared CancellationToken. The token is checked before and between operations. The service keeps one active token and one active scanner, protected by an asyncio lock, so two scans cannot overwrite progress or make cancellation ambiguous.",
            "ScanOrchestrator creates a short-lived scan record before collection begins. As each collector yields a ScanProgress event, the orchestrator upserts the current asset record and links an immutable JSON snapshot to the scan. The snapshot is the historical evidence; later asset updates cannot rewrite what an earlier scan observed.",
        ],
        "table": (
            ["Outcome", "Stored result"],
            [
                ["Completed", "Scan finalised, snapshots retained, optional collector artifacts stored, audit entry written."],
                ["Cancelled", "Scan marked cancelled and audited. It does not replace the current discovery or create false removals."],
                ["Failed", "Error retained on the scan and in audit history; exception returns to the route for an HTTP error."],
                ["Task cancellation", "Persistence happens before asyncio cancellation is propagated."],
            ],
        ),
        "bullets": [
            "Progress is shown as completed, total, and message; the scan itself never blocks the UI event loop.",
            "The active product route passes analyze_risk=False. Stored data is displayed as evidence rather than a severity score.",
        ],
        "code": "asset_id = repository.upsert_asset(site_id, progress.asset)\nrepository.link_scan_asset(scan.id, asset_id, observed_at, progress.asset)\n# The link stores the immutable observation for this scan.",
    },
    {
        "title": "SQLite model and the meaning of current state",
        "kicker": "Page 7 of 10 | persistence design",
        "paragraphs": [
            "SQLite is the local evidence store. Repository methods own connection setup, foreign-key enforcement, parameterized statements, schema initialization, and table writes. Dynamic values are always bound as SQL parameters. The database stores sites, scans, assets, ports, services, protocols, Wi-Fi access points, web observations, scan-to-asset snapshots, baselines, changes, and audit events.",
            "The assets table is a convenient current record with storage identity (site, IP). It is not the historical truth. scan_assets.snapshot_json preserves the full Asset that a collector yielded. The dashboard deliberately finds the latest completed discovery and uses its snapshots as current inventory. A later completed verification may overlay service information only for that discovered asset.",
        ],
        "table": (
            ["Data object", "Role in the system"],
            [
                ["sites", "Named scope for scans and observations."],
                ["scans", "Lifecycle state, target, level, timestamps, counts, cancellation, error."],
                ["assets / ports / services / protocols", "Latest observed records used for lookup and queries."],
                ["scan_assets", "Immutable JSON snapshot of an asset at a particular scan."],
                ["web_observations", "Raw bounded HTTP/TLS observations tied to the asset."],
                ["baselines / baseline_assets", "Named reference snapshot with active status."],
                ["changes", "Detected difference, analyst review state, note, and timestamps."],
                ["audit_log", "Operator action, result, duration, target, and error evidence."],
            ],
        ),
        "bullets": [
            "Current state excludes cancelled and failed discovery scans.",
            "Web metadata is filtered to observations collected after the current discovery finished, avoiding stale evidence.",
            "CSV export neutralizes spreadsheet formulas in user-controlled strings; JSON and CSV export the raw dashboard state rather than grades.",
        ],
    },
    {
        "title": "Baselines and change detection",
        "kicker": "Page 8 of 10 | comparing observations over time",
        "paragraphs": [
            "After reviewing a completed discovery, the operator can create a named baseline. The baseline stores asset and Wi-Fi snapshots, not live rows that may change later. Later completed scans are compared against the selected baseline and differences enter a review queue. Analysts can acknowledge a record and set its status to New, Investigating, Confirmed observation, Expected change, or Resolved with a note.",
            "Identity is deliberately conservative. The baseline comparison first reserves unique, case-normalized MAC matches. It only falls back to IP when there is no unique MAC match. This avoids a DHCP reassignment at one IP stealing the match from a device that moved to another IP. Missing MAC data is treated as missing evidence, not a MAC change.",
        ],
        "table": (
            ["Detected condition", "How it is derived"],
            [
                ["New device", "Current asset has no resolved baseline match."],
                ["Device removed", "Baseline asset has no resolved current match after a completed discovery."],
                ["MAC/IP changed", "A unique MAC match has a different IP, or both observed MAC values differ on an IP match."],
                ["New port", "A port exists in current verified evidence and not in the baseline snapshot."],
                ["Wi-Fi difference", "BSSID-backed access point observations differ from the baseline."],
            ],
        ),
        "bullets": [
            "A cancelled discovery never produces removal changes.",
            "A new discovery resets current service evidence until that host is checked again, preventing old ports from looking current.",
            "Persistent multi-factor identity reconciliation across MAC, IP, hostname, and confidence is a future improvement; hostname alone is not identity.",
        ],
    },
    {
        "title": "Safety model, security controls, and level 3",
        "kicker": "Page 9 of 10 | controls in code",
        "paragraphs": [
            "The local website binds to 127.0.0.1 and is designed for one local operator. It has browser security headers: no framing, no MIME sniffing, a same-origin content-security policy, no referrer, and no-store caching. The separate management_api package is a different engineering API and requires a bearer token with at least 32 characters; it is not the browser application's default backend.",
            "Safety is enforced in layers: route validation, service authorization checks, attached-target policy, scan level rules, low rate limits, subprocess timeouts, one active operation, cancellation, immutable evidence, and audit entries. This means the browser is not trusted as the only gate.",
        ],
        "table": (
            ["Level", "Current behavior"],
            [
                ["0 — Passive", "Modelled in the scanner contract; no packet capture implementation in the current website."],
                ["1 — Discovery", "Available: bounded Nmap host discovery on an authorized attached private subnet."],
                ["2 — Verify", "Available: confirmed one-host TCP check, web metadata inspection, and OS-managed Wi-Fi inventory."],
                ["3 — Active test", "Not available: planned protocol-specific read-only identity requests for one approved lab device."],
            ],
        ),
        "bullets": [
            "Read-only does not mean risk-free: a protocol request still transmits traffic and can affect fragile or legacy industrial equipment.",
            "Level 3 needs a lab-safe simulator or device, explicit authorization, one selected host, rate limits, cancellation, evidence logging, and test coverage before it can be enabled.",
            "No Raspberry Pi is required to develop the website or test levels 1–2 on an authorized local machine.",
        ],
    },
    {
        "title": "Running, testing, and extending the program",
        "kicker": "Page 10 of 10 | developer operations",
        "paragraphs": [
            "On Windows, run PowerShell from the repository root and execute .\\run.ps1. The launcher creates .venv if needed, installs the editable project, and starts the local website. The first launch needs Python 3.11+ and Nmap for discovery. The UI itself can open without Nmap, but live discovery and port checks require it.",
            "The project uses three complementary forms of validation. Unit and integration tests use controlled fixtures for repeatability. A live loopback website test starts the real FastAPI app and checks HTML, imported JavaScript modules, APIs, exports, headers, and authorization behavior. The local diagnostics command reads real local interfaces, starts a temporary 127.0.0.1 HTTP server, collects its response, and checks the server's open TCP port with Nmap.",
        ],
        "code": "cd D:\\proiecte\\ananas\n.\\run.ps1\n.\\.venv\\Scripts\\python.exe -m pytest -q\n.\\.venv\\Scripts\\python.exe -m pi_ot_probe.diagnostics.local_check",
        "table": (
            ["When changing…", "Start here"],
            [
                ["Page layout, labels, CSS", "web/static/index.html, css/styles.css, js/views/"],
                ["Button action or API error", "web/static/js/app.js and api.js"],
                ["Endpoint or input", "web/routes/ and web/schemas.py"],
                ["Scan workflow", "application/service.py and scanners/orchestrator.py"],
                ["Nmap / interface policy", "collectors/nmap.py, interfaces.py, targets.py"],
                ["Stored state / report", "database/, application/dashboard.py, reports/site.py"],
            ],
        ),
        "bullets": [
            "Use short comments for decisions and constraints, such as why a timeout exists or why a snapshot must stay immutable.",
            "Test behavior, not implementation names. The test suite currently contains 39 tests, including real local-loopback assembly coverage.",
            "Next programming priorities: stronger device identity reconciliation, safe OT lab simulators and level-3 checks, PDF export, and signed data-only update packs.",
        ],
    },
]


def build_markdown() -> str:
    output = ["# Pi-OT Security Probe - Technical Programming Guide", ""]
    for index, page in enumerate(PAGES, start=1):
        output.extend([f"## {index}. {page['title']}", ""])
        output.extend([paragraph + "\n" for paragraph in page.get("paragraphs", [])])
        if page.get("bullets"):
            output.extend([f"- {item}" for item in page["bullets"]])
            output.append("")
        if page.get("table"):
            headers, rows = page["table"]
            output.append("| " + " | ".join(headers) + " |")
            output.append("| " + " | ".join(["---"] * len(headers)) + " |")
            output.extend("| " + " | ".join(row) + " |" for row in rows)
            output.append("")
        if page.get("code"):
            output.extend(["```text", page["code"], "```", ""])
    output.extend([
        "## Source basis",
        "",
        "This document describes the implementation in `src/pi_ot_probe`, its tests, and the local launchers in this repository. It deliberately describes implemented behavior and separately labels deferred work.",
        "",
    ])
    return "\n".join(output)


def build_docx(output: Path) -> None:
    document = Document()
    style_document(document)
    for index, page in enumerate(PAGES, start=1):
        new_page(document, index, page["title"], page["kicker"])
        for paragraph in page.get("paragraphs", []):
            document.add_paragraph(paragraph)
        if page.get("bullets"):
            add_bullets(document, page["bullets"])
        if page.get("table"):
            headers, rows = page["table"]
            add_table(document, headers, rows)
        if page.get("code"):
            add_code(document, page["code"])
    document.save(output)


def ppt_textbox(slide, x, y, w, h, text: str, size: int = 18, color: str = INK,
                bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = ppt_rgb(color)


def ppt_bullets(slide, x: float, y: float, w: float, h: float, items: list[str], size: int = 16) -> None:
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = PptPt(size)
        paragraph.font.color.rgb = ppt_rgb(INK)
        paragraph.space_after = PptPt(10)
        paragraph.bullet = True


def ppt_box(slide, x: float, y: float, w: float, h: float, title: str, body: str,
            fill: str = WHITE, accent: str = TEAL) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ppt_rgb(fill)
    shape.line.color.rgb = ppt_rgb(LINE)
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(x), PptInches(y), PptInches(0.08), PptInches(h))
    strip.fill.solid()
    strip.fill.fore_color.rgb = ppt_rgb(accent)
    strip.line.fill.background()
    ppt_textbox(slide, x + 0.22, y + 0.17, w - 0.35, 0.32, title, 14, NAVY, True)
    ppt_textbox(slide, x + 0.22, y + 0.58, w - 0.35, h - 0.68, body, 11, MUTED)


def ppt_background(slide, number: int, title: str, subtitle: str = "") -> None:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = ppt_rgb(LIGHT)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(13.333), PptInches(0.18))
    band.fill.solid()
    band.fill.fore_color.rgb = ppt_rgb(TEAL)
    band.line.fill.background()
    ppt_textbox(slide, 0.6, 0.42, 11.8, 0.46, title, 26, NAVY, True)
    if subtitle:
        ppt_textbox(slide, 0.62, 0.95, 11.6, 0.32, subtitle, 11, MUTED)
    ppt_textbox(slide, 0.6, 7.13, 6.5, 0.18, "Pi-OT Security Probe | Programming and network evidence", 8, MUTED)
    ppt_textbox(slide, 12.0, 7.12, 0.7, 0.2, str(number), 8, MUTED, align=PP_ALIGN.RIGHT)


def add_slide(prs: Presentation, title: str, subtitle: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ppt_background(slide, len(prs.slides), title, subtitle)
    return slide


def build_pptx(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = ppt_rgb(NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(13.333), PptInches(0.2))
    accent.fill.solid(); accent.fill.fore_color.rgb = ppt_rgb(TEAL); accent.line.fill.background()
    ppt_textbox(slide, 0.8, 1.25, 10.5, 0.75, "Pi-OT Security Probe", 38, WHITE, True)
    ppt_textbox(slide, 0.82, 2.15, 10.4, 0.8, "A local, evidence-first network inventory console", 22, "CFE8F3")
    ppt_textbox(slide, 0.84, 3.2, 8.6, 1.0, "Programming architecture, safe discovery, evidence persistence, and development plan", 16, WHITE)
    ppt_box(slide, 8.9, 4.55, 3.45, 1.35, "What is implemented", "Local website • bounded discovery • service checks • web/TLS metadata • Wi-Fi inventory • baselines • exports", fill="173E5A", accent=ORANGE)
    ppt_textbox(slide, 0.85, 6.9, 6.5, 0.22, "Technical pitch deck | September 2026", 9, "CFE8F3")

    slide = add_slide(prs, "The problem: useful network evidence without overreach", "Operators need real observations, while developers need code they can safely change.")
    ppt_box(slide, 0.7, 1.55, 3.9, 4.75, "What is difficult", "Network inventories become unreliable when they overwrite history, mix current and stale service data, accept arbitrary targets, or turn weak signals into confident labels.", fill=WHITE, accent=ORANGE)
    ppt_box(slide, 4.75, 1.55, 3.9, 4.75, "Product response", "Collect only bounded evidence. Keep an immutable per-scan snapshot. Show missing evidence honestly. Separate scanning, persistence, routes, and views.", fill=WHITE, accent=TEAL)
    ppt_box(slide, 8.8, 1.55, 3.9, 4.75, "Outcome", "A local operator can see current devices, selected TCP evidence, raw web metadata, Wi-Fi observations, historical changes, and exports without risk scores.", fill=WHITE, accent=BLUE)

    slide = add_slide(prs, "What the current software delivers", "Runs on a normal Windows, Linux, or Kali machine; Raspberry Pi hardware is not required.")
    items = [
        ("Discover", "Find responsive hosts on one attached private /24-or-smaller network."),
        ("Verify", "Check 20 common TCP ports on one operator-selected discovered device."),
        ("Inspect", "Record a one-request HTTP/HTTPS observation for an observed web port."),
        ("Compare", "Create a baseline, review new/removed/IP/MAC/port/Wi-Fi changes, and add notes."),
        ("Export", "Download raw JSON and CSV evidence; retain scan and audit history."),
    ]
    for index, (name, description) in enumerate(items):
        x = 0.75 + (index % 3) * 4.18
        y = 1.55 + (index // 3) * 2.3
        ppt_box(slide, x, y, 3.85, 1.82, name, description, accent=[TEAL, BLUE, ORANGE, TEAL, BLUE][index])

    slide = add_slide(prs, "Architecture: responsibilities are separated", "Each folder owns one kind of decision, reducing the cost of a change.")
    layers = [
        ("web/", "Routes, schemas, browser assets", TEAL),
        ("application/", "Workflows, baselines, dashboard state", BLUE),
        ("collectors/", "Interfaces, target policy, Nmap, Wi-Fi, HTTP", ORANGE),
        ("scanners/ + database/", "Lifecycle, cancellation, audit, SQLite snapshots", TEAL),
        ("domain/", "Shared typed observations", BLUE),
    ]
    for index, (name, description, accent_color) in enumerate(layers):
        y = 1.4 + index * 1.0
        ppt_box(slide, 1.1, y, 11.0, 0.72, name, description, accent=accent_color)
    ppt_textbox(slide, 1.1, 6.6, 11, 0.25, "Deferred hardware and OT work lives in hardware/ and experimental/, outside the live website path.", 11, MUTED)

    slide = add_slide(prs, "Discovery flow: from a button click to historical evidence", "The same request path is covered by unit, integration, and live loopback tests.")
    flow = [
        ("Browser", "app.js + api.js"), ("Route", "schema + HTTP response"), ("Service", "authorize + lock"),
        ("Collector", "validate + Nmap"), ("Orchestrator", "progress + audit"), ("SQLite", "snapshot + state"),
    ]
    for index, (name, description) in enumerate(flow):
        x = 0.45 + index * 2.1
        ppt_box(slide, x, 2.35, 1.8, 1.3, name, description, accent=[TEAL, BLUE, ORANGE, TEAL, BLUE, ORANGE][index])
        if index < len(flow) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, PptInches(x + 1.83), PptInches(2.76), PptInches(0.22), PptInches(0.3))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = ppt_rgb(MUTED); arrow.line.fill.background()
    ppt_bullets(slide, 1.0, 4.55, 11.1, 1.7, [
        "Routes do not run shell commands; they validate input and call the application service.",
        "The service serializes operations and owns cancellation/progress state.",
        "The orchestrator stores an immutable scan snapshot before the browser reloads current state.",
    ], 15)

    slide = add_slide(prs, "Network discovery is bounded before Nmap starts", "Discovery first reads the local machine, then permits only a narrow set of targets.")
    rules = [
        ("Active local interface", "psutil first; Windows PowerShell or Linux iproute2 as fallback."),
        ("RFC1918 IPv4 only", "Public, loopback, and unrelated target ranges are rejected."),
        ("≤ 256 addresses", "A /24 or smaller subnet prevents accidental broad discovery."),
        ("Attached subnet", "Selected target must be contained by an active interface."),
        ("Authorization", "A server-side confirmation is required before the scan process starts."),
    ]
    for index, (name, description) in enumerate(rules):
        ppt_box(slide, 0.8, 1.45 + index * 0.95, 11.65, 0.69, name, description, accent=[TEAL, BLUE, ORANGE, TEAL, BLUE][index])

    slide = add_slide(prs, "Collectors have narrow, reviewable behavior", "Observations are retained as facts; the program does not claim vulnerabilities or device types.")
    collector_rows = [
        ("Host discovery", "Nmap -sn; rate/retry/host-time limits; parses hosts marked up."),
        ("Service check", "One current asset; TCP connect scan of top 20 ports; capped at 5 connections/sec."),
        ("HTTP/HTTPS", "One HEAD / request; no redirect following; redacts Set-Cookie; captures raw headers and TLS data."),
        ("Wi-Fi", "Normal OS-managed inventory through nmcli or netsh; no monitor mode or packet capture."),
    ]
    for index, (name, description) in enumerate(collector_rows):
        ppt_box(slide, 0.75 + (index % 2) * 6.05, 1.55 + (index // 2) * 2.25, 5.65, 1.72, name, description, accent=[TEAL, ORANGE, BLUE, TEAL][index])
    ppt_textbox(slide, 0.85, 6.25, 11.6, 0.35, "Excluded: exploit code, credential tests, NSE scripts, OS detection, broad scans, deauthentication, and live OT protocol checks.", 11, MUTED)

    slide = add_slide(prs, "Evidence stays meaningful over time", "The current device row is convenient; the per-scan snapshot is the historical truth.")
    ppt_box(slide, 0.8, 1.55, 3.7, 3.8, "Live collection", "Collector yields Asset and ScanProgress records. Each observation includes only fields actually returned by the tool or OS.", accent=TEAL)
    ppt_box(slide, 4.8, 1.55, 3.7, 3.8, "Persistence", "Repository upserts current rows and writes scan_assets.snapshot_json. Scan status and audit records are finalized for completed, cancelled, or failed outcomes.", accent=BLUE)
    ppt_box(slide, 8.8, 1.55, 3.7, 3.8, "Dashboard", "Latest completed discovery supplies current inventory. Later completed verification may overlay ports/services. Cancelled discovery cannot create false removals.", accent=ORANGE)
    ppt_textbox(slide, 1.0, 5.95, 11.25, 0.4, "Formula-safe CSV and raw JSON exports preserve observations without translating them into grades.", 14, NAVY, True, PP_ALIGN.CENTER)

    slide = add_slide(prs, "Baselines turn observations into an analyst review queue", "The comparison logic is conservative around changing IP addresses and missing data.")
    steps = [
        ("1. Create", "Store named asset and Wi-Fi snapshots after review."),
        ("2. Discover", "Run a later completed collection on the approved local target."),
        ("3. Match", "Reserve unique normalized MAC matches before IP fallback."),
        ("4. Review", "Record new/removed devices, MAC/IP changes, new ports, Wi-Fi differences."),
        ("5. Triage", "Keep analyst status, acknowledgement, and note alongside the observed difference."),
    ]
    for index, (name, description) in enumerate(steps):
        x = 0.6 + index * 2.5
        ppt_box(slide, x, 2.2, 2.2, 1.8, name, description, accent=[TEAL, BLUE, ORANGE, TEAL, BLUE][index])
    ppt_textbox(slide, 0.85, 5.05, 11.55, 0.85, "Missing MAC data is not treated as a MAC change. Hostname alone is not identity. Persistent multi-factor identity reconciliation is explicitly a future enhancement.", 15, MUTED, align=PP_ALIGN.CENTER)

    slide = add_slide(prs, "Safety model and level 3 roadmap", "The available features stop before industrial protocol requests.")
    levels = [
        ("Level 0", "Passive", "Contract only; packet capture is not implemented."),
        ("Level 1", "Discovery", "Available: bounded host discovery on one authorized attached private subnet."),
        ("Level 2", "Verify", "Available: one-host TCP check, web metadata, and managed Wi-Fi inventory."),
        ("Level 3", "Active test", "Disabled: planned protocol-specific read-only identity requests for one lab-approved device."),
    ]
    for index, (level, name, description) in enumerate(levels):
        ppt_box(slide, 0.75, 1.35 + index * 1.17, 11.75, 0.85, f"{level} — {name}", description, accent=[MUTED, TEAL, BLUE, ORANGE][index])
    ppt_textbox(slide, 0.85, 6.35, 11.6, 0.36, "Level 3 remains disabled because a read-only request still transmits traffic and needs lab-simulator coverage, explicit authorization, rate limits, cancellation, evidence logging, and protocol-specific tests.", 11, MUTED, align=PP_ALIGN.CENTER)

    slide = add_slide(prs, "Validation and development readiness", "The project is ready for software-focused iteration on an ordinary local machine.")
    ppt_box(slide, 0.75, 1.5, 3.8, 3.7, "Automated coverage", "39 tests cover safety gates, scan persistence, snapshots, baselines, reports, Nmap parsing, Wi-Fi parsing, and browser-route behavior.", accent=TEAL)
    ppt_box(slide, 4.78, 1.5, 3.8, 3.7, "Live local verification", "The diagnostics command reads actual interfaces, starts a temporary loopback HTTP server, captures its HTTP response, and checks its open TCP port with Nmap.", accent=BLUE)
    ppt_box(slide, 8.8, 1.5, 3.8, 3.7, "Change-friendly code", "The development guide maps each kind of edit to a file. Routes, application behavior, collectors, persistence, and views are independently testable.", accent=ORANGE)
    ppt_textbox(slide, 1.0, 5.95, 11.25, 0.38, "Run: .\\run.ps1  |  Test: .\\.venv\\Scripts\\python.exe -m pytest -q  |  Local evidence: python -m pi_ot_probe.diagnostics.local_check", 11, NAVY, True, PP_ALIGN.CENTER)

    slide = add_slide(prs, "Next programming milestones", "Build depth in controlled stages while preserving the current safety boundary.")
    roadmap = [
        ("Now", "Maintain the local website, improve views, collectors, reports, and tests."),
        ("Next", "Add stronger device identity reconciliation using MAC, IP, hostname, and explicit confidence."),
        ("Lab", "Implement protocol simulators and individually tested level-3 identity checks against a safe target."),
        ("Later", "Add PDF exports, data-only signed update packs with rollback, packet metadata capture, and remote management design."),
    ]
    for index, (phase, description) in enumerate(roadmap):
        ppt_box(slide, 0.9, 1.45 + index * 1.1, 11.45, 0.78, phase, description, accent=[TEAL, BLUE, ORANGE, TEAL][index])
    ppt_textbox(slide, 1.0, 6.35, 11.2, 0.35, "The project’s strongest current value is trustworthy, bounded local evidence with a codebase that a developer can understand and change safely.", 15, NAVY, True, PP_ALIGN.CENTER)

    prs.save(output)


def main() -> None:
    OUTPUT.mkdir(exist_ok=True)
    markdown = OUTPUT / "Pi-OT_Security_Probe_Technical_Programming_Guide.md"
    document = OUTPUT / "Pi-OT_Security_Probe_Technical_Programming_Guide.docx"
    deck = OUTPUT / "Pi-OT_Security_Probe_Pitch_Deck.pptx"
    markdown.write_text(build_markdown(), encoding="utf-8")
    build_docx(document)
    build_pptx(deck)
    print(f"Created {markdown}")
    print(f"Created {document}")
    print(f"Created {deck}")


if __name__ == "__main__":
    main()
